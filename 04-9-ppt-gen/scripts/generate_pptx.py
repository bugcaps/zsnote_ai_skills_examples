#!/usr/bin/env python3
"""Marp 마크다운을 .pptx로 변환합니다.

사용법:
    python generate_pptx.py deck.md deck.pptx
    python generate_pptx.py deck.md --check      # 변환하지 않고 구조만 검사

입력은 이 스킬이 만든 Marp 마크다운입니다. `---`로 슬라이드를 나누고, 첫 블록이
YAML 프론트매터(`marp: true`)이면 건너뜁니다. 각 슬라이드의 첫 제목(`#`/`##`)이
슬라이드 제목, `-` 목록이 본문이 됩니다.

python-pptx가 필요합니다. 없으면 변환하지 않고 설치 방법을 알린 뒤 종료 코드 3으로 끝납니다.
(--check 는 python-pptx 없이도 동작합니다.)
종료 코드: 0 = 성공, 1 = 입력 형식 오류, 2 = 사용법 오류, 3 = 의존성 없음
"""
import re
import sys

MAX_BULLETS = 4


def parse_marp(text):
    text = text.replace("\r\n", "\n")
    blocks = re.split(r"^---\s*$", text, flags=re.M)
    if blocks and not blocks[0].strip():
        blocks = blocks[1:]          # 파일 첫 줄의 `---` 앞 빈 블록
    if blocks and "marp:" in blocks[0] and not re.search(r"^#", blocks[0], flags=re.M):
        blocks = blocks[1:]          # YAML 프론트매터 블록
    slides = []
    for block in blocks:
        lines = [l.rstrip() for l in block.strip().split("\n") if l.strip()]
        if not lines:
            continue
        title, bullets = "", []
        for line in lines:
            heading = re.match(r"^#{1,3}\s+(.*)", line)
            if heading and not title:
                title = heading.group(1).strip()
            elif re.match(r"^[-*]\s+", line):
                bullets.append(re.sub(r"^[-*]\s+", "", line).strip())
        slides.append({"title": title, "bullets": bullets})
    return slides


def check(slides):
    problems = []
    for i, slide in enumerate(slides, 1):
        if not slide["title"]:
            problems.append(f"슬라이드 {i}: 제목(#)이 없습니다.")
        if len(slide["bullets"]) > MAX_BULLETS:
            problems.append(f"슬라이드 {i}: 불릿이 {len(slide['bullets'])}개입니다 (최대 {MAX_BULLETS}개).")
    for p in problems:
        print(f"FAIL {p}")
    print(f"요약: 슬라이드 {len(slides)}장, 문제 {len(problems)}건")
    return problems


def main():
    for stream in (sys.stdout, sys.stderr):  # Windows 콘솔에서 한글이 깨지지 않게 고정
        stream.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        return 2

    with open(args[0], encoding="utf-8") as f:
        slides = parse_marp(f.read())
    if not slides:
        print("오류: 슬라이드를 찾지 못했습니다. `---` 구분자와 제목을 확인하세요.", file=sys.stderr)
        return 1

    problems = check(slides)
    if len(args) > 1 and args[1] == "--check":
        return 1 if problems else 0
    if len(args) < 2:
        print("출력 파일 경로가 필요합니다. (구조만 보려면 --check)", file=sys.stderr)
        return 2
    if problems:
        print("구조 문제를 먼저 고친 뒤 다시 변환하세요.", file=sys.stderr)
        return 1

    try:
        from pptx import Presentation
    except ImportError:
        print("python-pptx가 없어 변환하지 않았습니다. `pip install python-pptx` 후 다시 실행하세요.", file=sys.stderr)
        print("마크다운만으로 충분하다면 Marp 뷰어나 `marp-cli`를 써도 됩니다.", file=sys.stderr)
        return 3

    prs = Presentation()
    for slide_data in slides:
        layout = prs.slide_layouts[1 if slide_data["bullets"] else 0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_data["title"]
        if slide_data["bullets"]:
            frame = slide.placeholders[1].text_frame
            frame.text = slide_data["bullets"][0]
            for bullet in slide_data["bullets"][1:]:
                frame.add_paragraph().text = bullet
    prs.save(args[1])
    print(f"저장 완료: {args[1]} (슬라이드 {len(slides)}장)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
